# -*- coding: utf-8 -*-
"""Convierte acta Markdown (formato PARA_FIRMAR) a .docx estilo S01/S02."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt


def strip_md_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    return text


def add_rich_paragraph(doc: Document, text: str, style: str = "Normal", bullet: bool = False):
    p = doc.add_paragraph(style="List Bullet" if bullet else style)
    if not text:
        return p
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            p.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`"):
            p.add_run(part[1:-1]).font.name = "Consolas"
        else:
            p.add_run(part)
    return p


def parse_table_lines(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        if not line.strip().startswith("|"):
            break
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.match(r"^:?-+:?$", c.replace(" ", "")) for c in cells):
            continue
        rows.append([strip_md_inline(c) for c in cells])
    return rows


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            val = row[j] if j < len(row) else ""
            table.rows[i].cells[j].text = val


def md_to_docx(md_path: Path, docx_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    i = 0
    first_h1 = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("# ") and not stripped.startswith("##"):
            text = stripped[2:].strip()
            if first_h1:
                doc.add_paragraph(text, style="Title")
                first_h1 = False
            else:
                doc.add_paragraph(text, style="Heading 1")
            i += 1
            continue

        if stripped.startswith("## "):
            doc.add_paragraph(stripped[3:].strip(), style="Heading 1")
            i += 1
            continue

        if stripped.startswith("### "):
            doc.add_paragraph(stripped[4:].strip(), style="Heading 2")
            i += 1
            continue

        if stripped.startswith("> "):
            add_rich_paragraph(doc, stripped[2:].strip())
            i += 1
            continue

        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, parse_table_lines(block))
            continue

        if stripped.startswith("- "):
            add_rich_paragraph(doc, stripped[2:].strip(), bullet=True)
            i += 1
            continue

        if stripped.startswith("_") and stripped.endswith("_") and len(stripped) > 2:
            p = doc.add_paragraph()
            p.add_run(stripped[1:-1]).italic = True
            i += 1
            continue

        add_rich_paragraph(doc, stripped)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def main():
    base = Path(__file__).resolve().parent.parent
    md = base / "DPAY3_ACTA_Retro_S03_PARA_FIRMAR.md"
    out1 = base / "DPAY3_ACTA_Retro_S03_PARA_FIRMAR.docx"
    out2 = (
        base.parent.parent
        / "Evidencias Proyecto"
        / "Evidencias de Documentacion"
        / "06_Evidencia_Retrospectiva"
        / "Sprint_03"
        / "DPAY3_ACTA_Retro_S03_PARA_FIRMAR.docx"
    )
    if len(sys.argv) >= 2:
        md = Path(sys.argv[1])
    if len(sys.argv) >= 3:
        out1 = Path(sys.argv[2])
        out2 = None
    md_to_docx(md, out1)
    print("Wrote", out1)
    if out2:
        md_to_docx(md, out2)
        print("Wrote", out2)


if __name__ == "__main__":
    main()
